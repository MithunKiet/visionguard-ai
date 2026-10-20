"""Durable report admission, restart recovery, fencing and export access."""
import asyncio
from datetime import datetime, timedelta, timezone
import io
import os
from unittest.mock import AsyncMock, Mock
from uuid import uuid4

import pytest
from sqlalchemy import func, select, update

from src.modules.reports.application import jobs
from src.modules.reports.application.services import ReportService
from src.modules.reports import runner
from src.shared.database.models import Alert, PPEViolation, Report, User
from src.shared.database.session import AsyncSessionFactory
from tests.test_review_integration import headers, setup as setup

pytestmark = pytest.mark.skipif(not os.environ.get("TEST_DATABASE_URL"), reason="Requires disposable PostgreSQL and Redis")
PAYLOAD = {"report_type": "alerts_summary", "format": "xlsx", "from_date": "2026-01-01T00:00:00Z", "to_date": "2026-01-02T00:00:00Z"}


async def enqueue(s, key=None, payload=None, caller="admin"):
    return await s.client.post("/api/v1/reports/generate", headers={**headers(s, caller),
        **({"Idempotency-Key": key} if key else {})}, json=payload or PAYLOAD)


async def details(s, report_id, caller="admin"):
    return await s.client.get(f"/api/v1/reports/{report_id}", headers=headers(s, caller))


async def test_request_only_persists_job_and_download_waits_for_completion(setup, monkeypatch):
    generate = Mock(side_effect=AssertionError("rendering must not run in request"))
    monkeypatch.setattr(runner, "build_xlsx", generate)
    response = await enqueue(setup)
    assert response.status_code == 202
    report = response.json()["data"]
    assert report["status"] == report["stage"] == "Queued" and report["attempts"] == 0
    assert "object_key" not in report and "lease_token" not in report
    assert (await details(setup, report["id"])).json()["data"]["status"] == "Queued"
    download = await setup.client.get(f"/api/v1/reports/{report['id']}/download", headers=headers(setup))
    assert download.status_code == 409
    generate.assert_not_called()


async def test_duplicate_requests_and_queue_capacity_are_serialized(setup):
    key = str(uuid4())
    responses = await asyncio.gather(enqueue(setup, key), enqueue(setup, key))
    assert all(r.status_code == 202 for r in responses)
    assert len({r.json()["data"]["id"] for r in responses}) == 1
    assert (await enqueue(setup, key, {**PAYLOAD, "format": "pdf"})).status_code == 409
    for _ in range(8):
        assert (await enqueue(setup)).status_code == 202
    last = await asyncio.gather(enqueue(setup), enqueue(setup))
    assert sorted(r.status_code for r in last) == [202, 429]
    # Duplicate admission remains safe even when the queue is full.
    assert (await enqueue(setup, key)).status_code == 202


@pytest.mark.parametrize("patch", [{"to_date": "2025-01-01T00:00:00Z"},
    {"to_date": "2028-01-01T00:00:00Z"}, {"from_date": "2026-01-01T00:00:00"}, {"format": "csv"}])
async def test_invalid_report_requests_are_rejected(setup, patch):
    assert (await enqueue(setup, payload={**PAYLOAD, **patch})).status_code == 422


async def test_reports_remain_tenant_scoped_and_admin_only(setup):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    for caller in ("manager", "viewer"):
        assert (await enqueue(setup, caller=caller)).status_code == 403
        for path in ("/reports", f"/reports/{report_id}", f"/reports/{report_id}/download"):
            assert (await setup.client.get("/api/v1" + path, headers=headers(setup, caller))).status_code == 403
        assert (await setup.client.post(f"/api/v1/reports/{report_id}/retry", headers=headers(setup, caller))).status_code == 403
    for suffix in ("", "/download"):
        assert (await setup.client.get(f"/api/v1/reports/{report_id}{suffix}", headers=headers(setup, "foreign"))).status_code == 404
    assert (await setup.client.post(f"/api/v1/reports/{report_id}/retry", headers=headers(setup, "foreign"))).status_code == 404
    assert (await setup.client.get("/api/v1/reports", headers=headers(setup, "foreign"))).json()["data"] == []


async def test_expired_claim_recovery_rejects_stale_progress_and_completion(setup):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    claims = await asyncio.gather(jobs.claim_report(report_id), jobs.claim_report(report_id))
    old = next(c for c in claims if c)
    assert sum(c is not None for c in claims) == 1
    async with AsyncSessionFactory() as db:
        await db.execute(update(Report).where(Report.public_id == report_id).values(lease_expires_at=func.now() - timedelta(seconds=1)))
        await db.commit()
    assert not await jobs.set_progress(*old, status="Completed", object_key="old")
    new = await jobs.claim_report(report_id)
    assert new[1] != old[1]
    assert not await jobs.set_progress(*old, stage="Uploading")
    await jobs.fail_attempt(*old, "STORAGE_UNAVAILABLE")
    assert (await details(setup, report_id)).json()["data"]["status"] == "Running"
    assert await jobs.set_progress(*new, status="Completed", stage="Completed", object_key="new", lease_token=None)


async def test_automatic_retry_is_bounded_and_manual_retry_is_idempotent(setup):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    for attempt in range(1, 4):
        claim = await jobs.claim_report(report_id)
        await jobs.fail_attempt(*claim, "STORAGE_UNAVAILABLE")
        row = (await details(setup, report_id)).json()["data"]
        assert row["attempts"] == attempt
        assert row["status"] == ("Failed" if attempt == 3 else "Retrying")
        assert await jobs.claim_report(report_id) is None
        async with AsyncSessionFactory() as db:
            await db.execute(update(Report).where(Report.public_id == report_id).values(next_attempt_at=func.now()))
            await db.commit()
    responses = await asyncio.gather(*[setup.client.post(f"/api/v1/reports/{report_id}/retry", headers=headers(setup)) for _ in range(2)])
    assert all(r.status_code == 202 and r.json()["data"]["status"] == "Queued" for r in responses)
    async with AsyncSessionFactory() as db:
        row = await db.scalar(select(Report).where(Report.public_id == report_id))
        assert row.attempt_limit == 6


async def test_final_crashed_attempt_becomes_failed_after_expiry(setup):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    await jobs.claim_report(report_id)
    async with AsyncSessionFactory() as db:
        await db.execute(update(Report).where(Report.public_id == report_id).values(
            attempts=3, lease_expires_at=func.now() - timedelta(seconds=1)))
        await db.commit()
    assert await jobs.claim_report(report_id) is None
    assert (await details(setup, report_id)).json()["data"]["status"] == "Failed"


@pytest.mark.parametrize("fmt,magic", [("pdf", b"%PDF"), ("xlsx", b"PK")])
async def test_real_rendering_records_stages_and_only_completed_report_downloads(setup, monkeypatch, fmt, magic):
    report_id = (await enqueue(setup, payload={**PAYLOAD, "format": fmt})).json()["data"]["id"]
    claim = await jobs.claim_report(report_id)
    uploaded = {}
    def put(bucket, key, stream, length, **kwargs):
        uploaded["content"] = stream.read()
        assert len(uploaded["content"]) == length
    storage = Mock(put_object=put)
    monkeypatch.setattr(runner, "storage_client", lambda: storage)
    stages = []
    original = runner.set_progress
    async def progress(*args, **kwargs):
        stages.append(kwargs["stage"])
        return await original(*args, **kwargs)
    monkeypatch.setattr(runner, "set_progress", progress)
    assert await runner.execute_report(*claim) == 0
    assert uploaded["content"].startswith(magic)
    assert stages == ["Rendering", "Uploading", "Completed"]
    row = (await details(setup, report_id)).json()["data"]
    assert row["status"] == "Completed" and row["row_count"] == 0
    storage.remove_object.assert_not_called()
    monkeypatch.setattr("src.modules.reports.application.services.get_presigned_url", lambda *a, **kw: "https://storage.example.invalid/report")
    assert (await setup.client.get(f"/api/v1/reports/{report_id}/download", headers=headers(setup))).status_code == 200
    assert (await setup.client.post(f"/api/v1/reports/{report_id}/retry", headers=headers(setup))).status_code == 409


async def test_revoked_requester_is_rechecked_before_work(setup, monkeypatch):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    claim = await jobs.claim_report(report_id)
    async with AsyncSessionFactory() as db:
        await db.execute(update(User).where(User.id == setup.users["admin"].id).values(status="Suspended"))
        await db.commit()
    storage = Mock()
    monkeypatch.setattr(runner, "storage_client", storage)
    assert await runner.execute_report(*claim) == 11
    storage.assert_not_called()


async def test_uncertain_completion_commit_never_deletes_potentially_referenced_file(setup, monkeypatch):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    claim = await jobs.claim_report(report_id)
    storage = Mock()
    monkeypatch.setattr(runner, "storage_client", lambda: storage)
    original = runner.set_progress
    async def ambiguous(*args, **kwargs):
        result = await original(*args, **kwargs)
        if kwargs.get("status") == "Completed":
            raise ConnectionError("lost commit reply")
        return result
    monkeypatch.setattr(runner, "set_progress", ambiguous)
    with pytest.raises(ConnectionError):
        await runner.execute_report(*claim)
    storage.remove_object.assert_not_called()
    assert (await details(setup, report_id)).json()["data"]["status"] == "Completed"


async def test_report_listing_has_stable_pagination(setup):
    async with AsyncSessionFactory() as db:
        db.add_all([Report(enterprise_id=setup.enterprise.id, generated_by=setup.users["admin"].id,
            report_type="alerts_summary", format="pdf", from_date=datetime.now(timezone.utc),
            to_date=datetime.now(timezone.utc), status="Completed", object_key="test-only") for _ in range(23)])
        await db.commit()
    first = (await setup.client.get("/api/v1/reports", headers=headers(setup))).json()
    second = (await setup.client.get("/api/v1/reports?page=2", headers=headers(setup))).json()
    assert first["meta"]["total"] == 23 and len(first["data"]) == 20 and len(second["data"]) == 3
    assert not {r["id"] for r in first["data"]} & {r["id"] for r in second["data"]}


async def test_worker_timeout_and_shutdown_leave_retryable_job(setup, monkeypatch):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    monkeypatch.setattr(jobs, "run_child", AsyncMock(return_value=124))
    assert await jobs.ReportWorker().process_one(report_id)
    assert (await details(setup, report_id)).json()["data"]["failure_code"] == "ATTEMPT_TIMEOUT"
    async with AsyncSessionFactory() as db:
        await db.execute(update(Report).where(Report.public_id == report_id).values(next_attempt_at=func.now()))
        await db.commit()
    monkeypatch.setattr(jobs, "run_child", AsyncMock(side_effect=asyncio.CancelledError))
    with pytest.raises(asyncio.CancelledError):
        await jobs.ReportWorker().process_one(report_id)
    assert (await details(setup, report_id)).json()["data"]["status"] == "Retrying"


async def test_real_subprocess_renders_uploads_and_commits_against_local_s3_fixture(setup, monkeypatch):
    """Real MinIO SDK HTTP transport; fixture is not a production MinIO server."""
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    import hashlib
    import threading
    from urllib.parse import urlsplit
    objects = {}
    class S3Fixture(BaseHTTPRequestHandler):
        def do_PUT(self):
            content = self.rfile.read(int(self.headers["Content-Length"]))
            objects[self.path] = content
            self.send_response(200)
            self.send_header("ETag", '"' + hashlib.md5(content).hexdigest() + '"')
            self.send_header("Content-Length", "0")
            self.end_headers()
        def do_GET(self):
            content = objects[urlsplit(self.path).path]
            self.send_response(200)
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        def log_message(self, *args):
            pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), S3Fixture)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    monkeypatch.setenv("MINIO_ENDPOINT", f"127.0.0.1:{server.server_port}")
    monkeypatch.setenv("MINIO_SECURE", "false")
    monkeypatch.setenv("MINIO_ACCESS_KEY", "test-only")
    monkeypatch.setenv("MINIO_SECRET_KEY", "test-only-secret")
    monkeypatch.setenv("MINIO_BUCKET_REPORTS", "report-job-test")
    try:
        report_id = (await enqueue(setup)).json()["data"]["id"]
        assert await jobs.ReportWorker().process_one(report_id)
        report = (await details(setup, report_id)).json()["data"]
        assert report["status"] == "Completed" and report["attempts"] == 1
        async with AsyncSessionFactory() as db:
            row = await db.scalar(select(Report).where(Report.public_id == report_id))
        content = objects["/report-job-test/" + row.object_key]
        import httpx
        from minio import Minio
        from src.core.settings import settings
        from src.shared.storage import minio_client
        monkeypatch.setattr(settings, "MINIO_BUCKET_REPORTS", "report-job-test")
        monkeypatch.setattr(minio_client, "_minio_public", Minio(f"127.0.0.1:{server.server_port}",
            access_key="test-only", secret_key="test-only-secret", secure=False, region="us-east-1"))
        response = await setup.client.get(f"/api/v1/reports/{report_id}/download", headers=headers(setup))
        url = response.json()["data"]["download_url"]
        assert "X-Amz-Signature=" in url
        async with httpx.AsyncClient() as client:
            assert (await client.get(url)).content == content
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content))
        assert workbook.active["A4"].value == "Alert #"
    finally:
        await asyncio.to_thread(server.shutdown)
        server.server_close()
        thread.join(timeout=2)


async def test_subprocess_timeout_reaps_native_work(monkeypatch):
    import sys
    real_spawn = asyncio.create_subprocess_exec
    captured = []
    async def slow_child(*args, **kwargs):
        child = await real_spawn(sys.executable, "-c", "import time; time.sleep(60)", **kwargs)
        captured.append(child)
        return child
    monkeypatch.setattr(jobs.asyncio, "create_subprocess_exec", slow_child)
    monkeypatch.setattr(jobs, "JOB_TIMEOUT_SECONDS", .05)
    assert await jobs.run_child(str(uuid4()), str(uuid4())) == 124
    assert captured[0].returncode is not None


async def test_row_limit_fails_instead_of_truncating_and_requester_is_rechecked_after_upload(setup, monkeypatch):
    from src.core.exceptions import VisionGuardException
    report_id = (await enqueue(setup)).json()["data"]["id"]
    claim = await jobs.claim_report(report_id)
    original = ReportService._alert_rows
    monkeypatch.setattr(ReportService, "_alert_rows", AsyncMock(side_effect=VisionGuardException("REPORT_TOO_LARGE", "limit", 422)))
    assert await runner.execute_report(*claim) == 10
    monkeypatch.setattr(ReportService, "_alert_rows", original)
    storage = Mock()
    monkeypatch.setattr(runner, "storage_client", lambda: storage)
    monkeypatch.setattr(runner, "authorized", AsyncMock(side_effect=[True, False]))
    assert await runner.execute_report(*claim) == 11
    storage.remove_object.assert_called_once()
    assert (await details(setup, report_id)).json()["data"]["status"] == "Running"


async def test_storage_failure_never_publishes_completed_report(setup, monkeypatch):
    report_id = (await enqueue(setup)).json()["data"]["id"]
    claim = await jobs.claim_report(report_id)
    monkeypatch.setattr(runner, "storage_client", lambda: Mock(put_object=Mock(side_effect=OSError("test outage"))))
    assert await runner.execute_report(*claim) == 13
    await jobs.fail_attempt(*claim, "STORAGE_UNAVAILABLE")
    assert (await details(setup, report_id)).json()["data"]["status"] == "Retrying"


@pytest.mark.parametrize("kind", ["alerts", "violations"])
async def test_export_rows_respect_tenant_dates_and_real_row_limit(setup, kind):
    from src.core.exceptions import VisionGuardException
    timestamp = datetime(2026, 1, 1, 10, 0, tzinfo=timezone(timedelta(hours=5, minutes=30)))
    async with AsyncSessionFactory() as db:
        for place, created in [(setup.places[0], timestamp), (setup.places[0], timestamp),
                               (setup.places[2], timestamp), (setup.places[0], timestamp - timedelta(days=2))]:
            values = dict(enterprise_id=place.factory.enterprise_id, zone_id=place.zone.id,
                          camera_id=place.camera.id, created_at=created)
            if kind == "alerts":
                db.add(Alert(**values, alert_number=uuid4().hex[:25], alert_type="PPE_TEST", severity="High"))
            else:
                db.add(PPEViolation(**values, violation_type="helmet_missing", confidence=.9))
        await db.commit()
        service = ReportService(db)
        extract = service._alert_rows if kind == "alerts" else service._violation_rows
        start, end = datetime(2026, 1, 1, tzinfo=timezone.utc), datetime(2026, 1, 2, tzinfo=timezone.utc)
        _, rows = await extract(setup.enterprise.id, start, end, limit=2)
        assert len(rows) == 2 and all("2026-01-01 04:30" in row for row in rows)
        with pytest.raises(VisionGuardException) as failure:
            await extract(setup.enterprise.id, start, end, limit=1)
        assert failure.value.code == "REPORT_TOO_LARGE" and "row limit" in failure.value.message


def test_excel_cells_remain_literal_and_pdf_escapes_branding():
    from openpyxl import load_workbook
    from src.modules.reports.infrastructure.generators import build_pdf, build_xlsx
    workbook = load_workbook(io.BytesIO(build_xlsx("Report", "=Hostile", "Test", ["Zone"], [["=1+1"], ["@SUM(1)"]])))
    assert workbook.active["A1"].data_type == "s" and workbook.active["A5"].data_type == "s"
    assert workbook.active["A5"].value == "=1+1"
    assert build_pdf("Report", "Factory <unclosed>", "Test", ["Zone"], [["Line A"]]).startswith(b"%PDF")
